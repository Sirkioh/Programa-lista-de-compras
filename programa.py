import openpyxl
from openpyxl.styles import PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import requests
import re
import json

compras = {}

supermercados = [
    "California",
    "Yaguar",
    "Multiexpress",
    "Changomas",
    "HiperDelPollo",
    "Maxiconsumo",
    "Vital",
    "La Anónima",
    "Diarco",
    "Makro"
]

def crearExcel():
    global compras
    libro = openpyxl.Workbook()
    hoja = libro.active
    tituloExcel = ["Producto", "Categoria", "Q"] + supermercados
    for n, v in enumerate(tituloExcel, start=1):
        hoja[f"{get_column_letter(n)}1"] = v
    hoja.freeze_panes = "A2"
    
    azulOscuro = PatternFill("solid", fgColor="3F6F8F")
    azulClaro = PatternFill("solid", fgColor="A9C5D8")
    verdePrimero = PatternFill("solid", fgColor="D4E6D1")
    verdeSegundo = PatternFill("solid", fgColor="E5F0E3")
    verdePrecioMasBarato = PatternFill("solid", fgColor="118A25")

    preciosSupers = {}
    for p,v in compras.items():    
        nombreProducto = v["nombreNormalizado"].capitalize()
        valores = [nombreProducto,]
        precioMasBarato = []

        valores.append(v["categoria"])
        valores.append(v["cantidad"])
        for s in supermercados:
            p = v["precios"].get(s)
            precioTotal = None if p == None else p * v["cantidad"]
            valores.append(precioTotal)
            if precioTotal is not None: precioMasBarato.append(precioTotal)
        precioMasBarato = min(precioMasBarato)
        preciosSupers[nombreProducto] = precioMasBarato
        hoja.append(valores)
        filaActual = hoja.max_row
        if filaActual % 2 == 0:
            for celda in hoja[filaActual]:
                celda.fill = verdePrimero
        else:
            for celda in hoja[filaActual]:
                celda.fill = verdeSegundo 
                
        for celda in hoja[filaActual]:
            if celda.value == precioMasBarato:
                celda.fill = verdePrecioMasBarato
            if isinstance(celda.value, (int,float)): celda.number_format = "$#.##0"
                
                
    
    hoja.append(["Total $ mas barato", sum(preciosSupers.values())])
    hoja[f"B{hoja.max_row}"].number_format = "$#.##0"
            
    
    for celda in hoja["A"]:
        celda.fill = azulClaro
    for celda in hoja[1]:
            celda.fill = azulOscuro
            
    for i in range(2, 14): hoja.column_dimensions[get_column_letter(i)].width = 15
    hoja.column_dimensions["A"].width = 40
    for celda in hoja[1]: celda.alignment = Alignment(horizontal="center")
    for i in range(3, 14):
        for celda in hoja[get_column_letter(i)]: celda.alignment = Alignment(horizontal="center")
    
    bordeGrueso = Side(style="thick")
    bordeFino = Side(style="thin")
    ultimaFila = hoja.max_row

    for fila in hoja[f"A1:M{ultimaFila}"]:
        for celda in fila:
            celda.border = Border(
                top=bordeGrueso if celda.row == 1 else bordeFino,
                bottom=bordeGrueso if celda.row == len(compras) + 2 else bordeFino,
                left=bordeGrueso if celda.column == 1 else bordeFino,
                right=bordeGrueso if celda.column == 13 else bordeFino
            )
    for celda in hoja[hoja.max_row]:
        celda.fill = azulClaro
    
    libro.save("compras.xlsx")
    compras = {}

def buscarProductos(producto): 
    parametros = {"q": producto,"page": 1,"limit": 24}

    respuestaBusqueda = requests.get("https://ahorroposadas.com/api/search", params=parametros, verify=False)
    datosBusqueda = respuestaBusqueda.json()
    productos = datosBusqueda["products"]
    
    productosNombres = {}
    for p in productos:
        productosNombres[p["normalizedName"]] = p["relevanceScore"]
    
    prodcutosMayoresScore = sorted(productosNombres, key=productosNombres.get, reverse=True)[:10]
    return [prodcutosMayoresScore, productos]

def buscarPrecios(productoElegido, productos, cantidad, producto):
    global compras
    for p in productos:
        if p["normalizedName"] == productoElegido:
            productoCategoria = p["category"]
            break
    
    datosProducto = requests.get(f"https://ahorroposadas.com/producto/{productoElegido}", verify=False).text
    match = re.search(r'const products = (\[.*?\]);', datosProducto, re.DOTALL)
    
    products = json.loads(match.group(1))
    datosProducto = {
        "precios" : {},
        "nombreNormalizado" : productoElegido,
        "categoria" : productoCategoria
    }
    for p in products:
        datosProducto["precios"][p["supermarket"]] = p["price"]
    

    compras[producto] = {"cantidad": int(cantidad)}
    compras[producto].update(datosProducto)
