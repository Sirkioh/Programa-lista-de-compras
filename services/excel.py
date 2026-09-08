import openpyxl
from openpyxl.styles import PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class xl:
    def __init__(self):
        self.compras = {}
        self.supermercados = [
            "California",
            "Yaguar",
            "Multiexpress",
            "Changomas",
            "HiperDelPollo",
            "Maxiconsumo",
            "Vital",
            "La Anónima",
            "Diarco",
            "Makro"
        ]
        self.libro = openpyxl.Workbook()
        self.hoja = self.libro.active

        self.azulOscuro = PatternFill("solid", fgColor="3F6F8F")
        self.azulClaro = PatternFill("solid", fgColor="A9C5D8")
        self.verdePrimero = PatternFill("solid", fgColor="D4E6D1")
        self.verdeSegundo = PatternFill("solid", fgColor="E5F0E3")
        self.verdePrecioMasBarato = PatternFill("solid", fgColor="118A25")

    def crearExcel(self):
        tituloExcel = ["Producto", "Categoria", "Q"] + self.supermercados
        for n, v in enumerate(tituloExcel, start=1):
            self.hoja[f"{get_column_letter(n)}1"] = v
        self.hoja.freeze_panes = "A2"

        preciosSupers = {}
        for p,v in self.compras.items():
            nombreProducto = v["nombreNormalizado"].capitalize()
            valores = [nombreProducto,]
            precioMasBarato = []

            valores.append(v["categoria"])
            valores.append(v["cantidad"])
            for s in self.supermercados:
                p = v["precios"].get(s)
                precioTotal = None if p == None else p * v["cantidad"]
                valores.append(precioTotal)
                if precioTotal is not None: precioMasBarato.append(precioTotal)
            precioMasBarato = min(precioMasBarato)
            preciosSupers[nombreProducto] = precioMasBarato
            self.hoja.append(valores)
            filaActual = self.hoja.max_row
            if filaActual % 2 == 0:
                for celda in self.hoja[filaActual]:
                    celda.fill = self.verdePrimero
            else:
                for celda in self.hoja[filaActual]:
                    celda.fill = self.verdeSegundo 
                    
            for celda in self.hoja[filaActual]:
                if celda.value == precioMasBarato:
                    celda.fill = self.verdePrecioMasBarato
                if isinstance(celda.value, (int,float)): celda.number_format = "$#.##0"
                    
                    
        
        self.hoja.append(["Total $ mas barato", sum(preciosSupers.values())])
        self.hoja[f"B{self.hoja.max_row}"].number_format = "$#.##0"
                
        
        for celda in self.hoja["A"]:
            celda.fill = self.azulClaro
        for celda in self.hoja[1]:
                celda.fill = self.azulOscuro
                
        for i in range(2, 14): self.hoja.column_dimensions[get_column_letter(i)].width = 15
        self.hoja.column_dimensions["A"].width = 40
        for celda in self.hoja[1]: celda.alignment = Alignment(horizontal="center")
        for i in range(3, 14):
            for celda in self.hoja[get_column_letter(i)]: celda.alignment = Alignment(horizontal="center")
        
        bordeGrueso = Side(style="thick")
        bordeFino = Side(style="thin")
        ultimaFila = self.hoja.max_row

        for fila in self.hoja[f"A1:M{ultimaFila}"]:
            for celda in fila:
                celda.border = Border(
                    top=bordeGrueso if celda.row == 1 else bordeFino,
                    bottom=bordeGrueso if celda.row == len(self.compras) + 2 else bordeFino,
                    left=bordeGrueso if celda.column == 1 else bordeFino,
                    right=bordeGrueso if celda.column == 13 else bordeFino
                )
        for celda in self.hoja[self.hoja.max_row]:
            celda.fill = self.azulClaro
        
        self.libro.save("compras.xlsx")
        self.compras = {}
